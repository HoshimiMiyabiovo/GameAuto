from __future__ import annotations

import json
import logging
from collections import Counter, defaultdict
from datetime import datetime, timezone, timedelta
from typing import Any, Dict, List, Optional, Tuple

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

CST = timezone(timedelta(hours=8))

SHEET_POOL_MAP = {
    "角色活动唤取": "限定角色池",
    "武器活动唤取": "限定武器池",
    "角色常驻唤取": "常驻角色池",
    "武器常驻唤取": "常驻武器池",
    "新手唤取": "新手池",
    "新手自选唤取": "新手自选池",
    "新手自选唤取（感恩定向唤取）": "感恩定向唤取",
    "角色新旅唤取": "角色新旅池",
    "武器新旅唤取": "武器新旅池",
    "角色联动唤取": "角色联动池",
    "武器联动唤取": "武器联动池",
    "角色忆旅唤取": "角色忆旅池",
    "武器忆旅唤取": "武器忆旅池",
}

RARITY_LABEL = {5: "五星", 4: "四星", 3: "三星"}
RARITY_COLORS = {5: "#d4a017", 4: "#8e44ad", 3: "#2980b9"}

SHEET_HEADERS = ["名称", "类型", "星级", "时间", "数量", "资源ID", "卡池"]

_logger = logging.getLogger("WuwaGacha")


class WuwaGachaRecord:
    __slots__ = (
        "name", "item_type", "rarity", "time_str", "quantity",
        "resource_id", "pool_name", "pool_display", "sheet_name", "gacha_time_str",
    )

    def __init__(
        self,
        name: str,
        item_type: str,
        rarity: int,
        time_str: str,
        quantity: int,
        resource_id: str,
        pool_name: str,
        pool_display: str,
        sheet_name: str,
        gacha_time_str: str,
    ):
        self.name = name
        self.item_type = item_type
        self.rarity = rarity
        self.time_str = time_str
        self.quantity = quantity
        self.resource_id = resource_id
        self.pool_name = pool_name
        self.pool_display = pool_display
        self.sheet_name = sheet_name
        self.gacha_time_str = gacha_time_str

    def to_dict(self) -> Dict[str, Any]:
        return {
            "name": self.name,
            "item_type": self.item_type,
            "rarity": self.rarity,
            "rarity_label": RARITY_LABEL.get(self.rarity, f"{self.rarity}星"),
            "time": self.time_str,
            "quantity": self.quantity,
            "resource_id": self.resource_id,
            "pool_name": self.pool_name,
            "pool_display": self.pool_display,
            "sheet_name": self.sheet_name,
        }


def _parse_xlsx_time_to_iso(raw_time: object) -> Tuple[str, str]:
    if raw_time is None:
        dt = datetime(2000, 1, 1, tzinfo=CST)
    elif isinstance(raw_time, datetime):
        dt = raw_time
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=CST)
    elif isinstance(raw_time, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S"):
            try:
                dt = datetime.strptime(str(raw_time).strip(), fmt).replace(tzinfo=CST)
                break
            except ValueError:
                continue
        else:
            dt = datetime(2000, 1, 1, tzinfo=CST)
    else:
        dt = datetime(2000, 1, 1, tzinfo=CST)

    return dt.strftime("%Y-%m-%d %H:%M:%S"), dt.isoformat()


def load_wuwa_gacha(file_path: str) -> Dict[str, List[WuwaGachaRecord]]:
    if openpyxl is None:
        raise ImportError("缺少 openpyxl 库，请 pip install openpyxl")

    wb = openpyxl.load_workbook(file_path, read_only=True)
    results: Dict[str, List[WuwaGachaRecord]] = {}

    for sheet_name in wb.sheetnames:
        pool_display = SHEET_POOL_MAP.get(sheet_name)
        if pool_display is None:
            continue

        ws = wb[sheet_name]
        headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
        if headers[:7] != SHEET_HEADERS:
            _logger.warning(f"Sheet {sheet_name} headers mismatch: {headers}")
            continue

        records: List[WuwaGachaRecord] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(v is None for v in row):
                continue
            name = str(row[0] or "")
            item_type = str(row[1] or "角色")
            try:
                rarity = int(float(str(row[2] or 3)))
            except (ValueError, TypeError):
                rarity = 3
            quantity = 1
            if row[4] is not None:
                try:
                    quantity = int(float(str(row[4])))
                except (ValueError, TypeError):
                    quantity = 1
            resource_id = str(row[5] or "")
            pool_name = str(row[6] or sheet_name)
            time_str, gacha_time_str = _parse_xlsx_time_to_iso(row[3])

            records.append(WuwaGachaRecord(
                name=name,
                item_type=item_type,
                rarity=rarity,
                time_str=time_str,
                quantity=quantity,
                resource_id=resource_id,
                pool_name=pool_name,
                pool_display=pool_display,
                sheet_name=sheet_name,
                gacha_time_str=gacha_time_str,
            ))

        if records:
            results[sheet_name] = records
            _logger.info(f"Sheet {sheet_name} ({pool_display}): {len(records)} records")

    wb.close()
    return results


def load_all_wuwa_records(file_path: str) -> List[WuwaGachaRecord]:
    grouped = load_wuwa_gacha(file_path)
    all_records: List[WuwaGachaRecord] = []
    for sheet_records in grouped.values():
        all_records.extend(sheet_records)
    all_records.sort(key=lambda r: r.gacha_time_str)
    return all_records


def _parse_pool_json_time_to_iso(raw_time: str) -> Tuple[str, str]:
    """解析 WutheringWavesTool pool.json 的时间字段（"2026-07-30 20:17:42"）。

    Returns:
        (display_str, iso_str): 例如 ("2026-07-30 20:17:42", "2026-07-30T20:17:42+08:00")
    """
    if not raw_time:
        dt = datetime(2000, 1, 1, tzinfo=CST)
    else:
        dt = None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S"):
            try:
                dt = datetime.strptime(str(raw_time).strip(), fmt).replace(tzinfo=CST)
                break
            except ValueError:
                continue
        if dt is None:
            dt = datetime(2000, 1, 1, tzinfo=CST)
    return dt.strftime("%Y-%m-%d %H:%M:%S"), dt.isoformat()


def load_wuwa_pool_json(file_paths) -> List[WuwaGachaRecord]:
    """从 WutheringWavesTool 的 pool.json 加载、合并、去重抽卡记录。

    支持单个路径（str）或多个路径（List[str]）。多文件会合并后按业务键
    (gacha_time_str, resource_id, cardPoolType) 去重，按时间升序排列。

    Args:
        file_paths: pool.json 路径字符串或路径列表

    Returns:
        合并去重后的 WuwaGachaRecord 列表
    """
    if isinstance(file_paths, str):
        file_paths = [file_paths]
    elif not isinstance(file_paths, (list, tuple)):
        raise TypeError(f"file_paths 必须是 str 或 List[str]: {type(file_paths)}")

    seen: Dict[Tuple[str, str, str], WuwaGachaRecord] = {}

    for file_path in file_paths:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except (json.JSONDecodeError, OSError) as exc:
            _logger.warning(f"读取 pool.json 失败 {file_path}: {exc}")
            continue

        if not isinstance(data, dict):
            _logger.warning(f"pool.json 顶层不是 dict: {file_path}")
            continue

        per_file_count = 0
        for sheet_name, items in data.items():
            pool_display = SHEET_POOL_MAP.get(sheet_name)
            if pool_display is None:
                _logger.info(f"跳过未知卡池 sheet: {sheet_name} (file={file_path})")
                continue
            if not isinstance(items, list):
                continue

            for item in items:
                if not isinstance(item, dict):
                    continue
                try:
                    resource_id = str(item.get("resourceId", ""))
                    time_str, gacha_time_str = _parse_pool_json_time_to_iso(
                        str(item.get("time", ""))
                    )
                    pool_name = str(item.get("cardPoolType", sheet_name))
                    name = str(item.get("name", ""))
                    item_type = str(item.get("resourceType", "角色"))
                    try:
                        rarity = int(item.get("qualityLevel", 3))
                    except (ValueError, TypeError):
                        rarity = 3
                    try:
                        quantity = int(item.get("count", 1))
                    except (ValueError, TypeError):
                        quantity = 1

                    # 业务键去重：(gacha_time_str, resource_id, pool_name)
                    key = (gacha_time_str, resource_id, pool_name)
                    if key in seen:
                        continue

                    seen[key] = WuwaGachaRecord(
                        name=name,
                        item_type=item_type,
                        rarity=rarity,
                        time_str=time_str,
                        quantity=quantity,
                        resource_id=resource_id,
                        pool_name=pool_name,
                        pool_display=pool_display,
                        sheet_name=sheet_name,
                        gacha_time_str=gacha_time_str,
                    )
                    per_file_count += 1
                except Exception as exc:
                    _logger.warning(f"解析记录失败 item={item} err={exc}")
                    continue

        _logger.info(
            f"加载 pool.json: {file_path}, 本文件新增 {per_file_count} 条, "
            f"累计去重后 {len(seen)} 条"
        )

    records = list(seen.values())
    records.sort(key=lambda r: r.gacha_time_str)
    return records


def wuwa_pool_json_to_db_records(
    records: List[WuwaGachaRecord], uid: str
) -> List[Dict[str, Any]]:
    """把 WuwaGachaRecord 列表转为 db 记录。

    与 wuwa_to_db_records 兼容，但 record_id 用业务键
    `wuwa_{uid}_{time_compact}_{resource_id}`，保证同一抽卡在多次导入里
    record_id 一致，被 upsert 自动去重。
    """
    results: List[Dict[str, Any]] = []
    for r in records:
        time_compact = (
            r.gacha_time_str
            .replace("-", "")
            .replace(":", "")
            .replace("+", "")
            .replace("T", "T")
        )
        record_id = f"wuwa_{uid}_{time_compact}_{r.resource_id}"
        raw = json.dumps(r.to_dict(), ensure_ascii=False)
        results.append({
            "record_id": record_id,
            "gacha_type": r.pool_name,
            "uigf_gacha_type": "",
            "item_id": r.resource_id,
            "name": r.name,
            "item_type": r.item_type,
            "rank_type": str(r.rarity),
            "count": r.quantity,
            "time": r.gacha_time_str,
            "lang": "zh-cn",
            "timezone": "UTC+8",
            "raw_json": raw,
        })
    return results


def build_wuwa_overview(records: List[WuwaGachaRecord]) -> Dict[str, Any]:
    if not records:
        return {"total_records": 0, "five_star_count": 0, "four_star_count": 0, "three_star_count": 0,
                "pool_counts": {}, "earliest_time": "", "latest_time": "", "top_five_star": ("-", 0)}

    pool_counter = Counter(r.pool_display for r in records)
    rarity_counter = Counter(str(r.rarity) for r in records)
    five_star_names = Counter(r.name for r in records if r.rarity == 5)

    return {
        "total_records": len(records),
        "five_star_count": rarity_counter.get("5", 0),
        "four_star_count": rarity_counter.get("4", 0),
        "three_star_count": rarity_counter.get("3", 0),
        "pool_counts": {name: count for name, count in sorted(pool_counter.items())},
        "earliest_time": records[0].time_str,
        "latest_time": records[-1].time_str,
        "top_five_star": five_star_names.most_common(1)[0] if five_star_names else ("-", 0),
    }


def build_wuwa_pool_analysis(records: List[WuwaGachaRecord]) -> Dict[str, Any]:
    pools: Dict[str, Dict[str, Any]] = {}
    pity_counters: Dict[str, int] = defaultdict(int)

    for r in records:
        pool_display = r.pool_display
        if pool_display not in pools:
            pools[pool_display] = {
                "total": 0,
                "five_star_count": 0, "four_star_count": 0, "three_star_count": 0,
                "sub_pools": defaultdict(int),
                "earliest_time": r.time_str,
                "latest_time": r.time_str,
                "five_star_list": [],
                "character_count": 0, "weapon_count": 0, "tool_count": 0,
            }
        p = pools[pool_display]
        p["total"] += r.quantity
        pity_counters[pool_display] += r.quantity
        rarity = r.rarity
        if rarity == 5:
            p["five_star_count"] += r.quantity
            p["five_star_list"].append({
                "name": r.name, "time": r.time_str, "pool_name": r.pool_name,
                "type": r.item_type,
                "pity": pity_counters[pool_display],
            })
            pity_counters[pool_display] = 0
        elif rarity == 4:
            p["four_star_count"] += r.quantity
        elif rarity == 3:
            p["three_star_count"] += r.quantity
        p["sub_pools"][r.pool_name] += r.quantity
        if r.item_type == "角色":
            p["character_count"] += r.quantity
        elif r.item_type == "武器":
            p["weapon_count"] += r.quantity
        elif r.item_type == "道具":
            p["tool_count"] += r.quantity
        p["latest_time"] = r.time_str

    for pool_display in pools:
        p = pools[pool_display]
        if p["total"] > 0:
            p["five_star_pct"] = round(p["five_star_count"] / p["total"] * 100, 2)
            p["four_star_pct"] = round(p["four_star_count"] / p["total"] * 100, 2)
            p["three_star_pct"] = round(p["three_star_count"] / p["total"] * 100, 2)
        else:
            p["five_star_pct"] = p["four_star_pct"] = p["three_star_pct"] = 0
        p["sub_pools"] = dict(sorted(p["sub_pools"].items(), key=lambda x: -x[1]))

    return pools


def build_wuwa_five_star_history(records: List[WuwaGachaRecord], limit: int = 120) -> List[Dict[str, Any]]:
    pity_counters: Dict[str, int] = defaultdict(int)
    history: List[Dict[str, Any]] = []

    for r in records:
        group = r.pool_display
        pity_counters[group] += 1
        if r.rarity != 5:
            continue
        history.append({
            "time": r.time_str,
            "name": r.name,
            "pool": r.pool_display,
            "pool_name": r.pool_name,
            "type": r.item_type,
            "pity": pity_counters[group],
        })
        pity_counters[group] = 0

    return list(reversed(history[-limit:]))


def build_wuwa_pity_analysis(records: List[WuwaGachaRecord]) -> Dict[str, Any]:
    pity_counters: Dict[str, int] = defaultdict(int)
    pity_history: Dict[str, List[int]] = defaultdict(list)

    for r in records:
        group = r.pool_display
        pity_counters[group] += 1
        if r.rarity == 5:
            pity_history[group].append(pity_counters[group])
            pity_counters[group] = 0

    result: Dict[str, Any] = {}
    all_pool_keys = set(pity_history.keys()) | set(pity_counters.keys())
    for group in all_pool_keys:
        pities = pity_history.get(group, [])
        if not pities:
            result[group] = {"count": 0, "avg_pity": "-", "min_pity": "-", "max_pity": "-", "pities": [],
                             "current_pity_no_5s": pity_counters.get(group, 0), "distribution": {}}
            continue
        result[group] = {
            "count": len(pities),
            "avg_pity": round(sum(pities) / len(pities), 1),
            "min_pity": min(pities),
            "max_pity": max(pities),
            "current_pity_no_5s": pity_counters.get(group, 0),
            "pities": pities,
            "distribution": dict(Counter(pities).most_common()),
        }
    return result


def build_wuwa_character_stats(records: List[WuwaGachaRecord]) -> List[Dict[str, Any]]:
    grouped: Dict[tuple, List[WuwaGachaRecord]] = defaultdict(list)
    for r in records:
        if r.item_type != "角色":
            continue
        grouped[(r.name, r.rarity)].append(r)

    rows: List[Dict[str, Any]] = []
    for (name, rarity), items in grouped.items():
        rows.append({
            "name": name,
            "rarity": rarity,
            "rarity_label": RARITY_LABEL.get(rarity, f"{rarity}星"),
            "copies": len(items),
            "first_time": items[0].time_str,
            "last_time": items[-1].time_str,
        })
    rows.sort(key=lambda r: (-r["rarity"], -r["copies"], r["last_time"], r["name"]))
    return rows


def wuwa_to_db_records(records: List[WuwaGachaRecord], uid: str) -> List[Dict[str, Any]]:
    results: List[Dict[str, Any]] = []
    for i, r in enumerate(records):
        raw = json.dumps(r.to_dict(), ensure_ascii=False)
        results.append({
            "record_id": f"wuwa_{uid}_{r.resource_id}_{i}",
            "gacha_type": r.pool_name,
            "uigf_gacha_type": "",
            "item_id": r.resource_id,
            "name": r.name,
            "item_type": r.item_type,
            "rank_type": str(r.rarity),
            "count": r.quantity,
            "time": r.gacha_time_str,
            "lang": "zh-cn",
            "timezone": "UTC+8",
            "raw_json": raw,
        })
    return results


def db_to_wuwa_records(db_rows: List[Dict[str, Any]]) -> List[WuwaGachaRecord]:
    results: List[WuwaGachaRecord] = []
    for row in db_rows:
        try:
            raw = json.loads(row.get("raw_json", "{}"))
        except (json.JSONDecodeError, TypeError):
            raw = {}
        results.append(WuwaGachaRecord(
            name=str(raw.get("name", row.get("name", ""))),
            item_type=str(raw.get("item_type", row.get("item_type", ""))),
            rarity=int(raw.get("rarity", row.get("rank_type", 3))),
            time_str=str(raw.get("time", row.get("time", ""))),
            quantity=int(raw.get("quantity", row.get("count", 1))),
            resource_id=str(raw.get("resource_id", row.get("item_id", ""))),
            pool_name=str(raw.get("pool_name", row.get("gacha_type", ""))),
            pool_display=str(raw.get("pool_display", "")),
            sheet_name=str(raw.get("sheet_name", "")),
            gacha_time_str=str(raw.get("time", row.get("time", ""))),
        ))
    return results


def build_wuwa_weapon_stats(records: List[WuwaGachaRecord]) -> List[Dict[str, Any]]:
    grouped: Dict[tuple, List[WuwaGachaRecord]] = defaultdict(list)
    for r in records:
        if r.item_type != "武器":
            continue
        grouped[(r.name, r.rarity)].append(r)

    rows: List[Dict[str, Any]] = []
    for (name, rarity), items in grouped.items():
        rows.append({
            "name": name,
            "rarity": rarity,
            "rarity_label": RARITY_LABEL.get(rarity, f"{rarity}星"),
            "copies": len(items),
            "first_time": items[0].time_str,
            "last_time": items[-1].time_str,
        })
    rows.sort(key=lambda r: (-r["rarity"], -r["copies"], r["last_time"], r["name"]))
    return rows